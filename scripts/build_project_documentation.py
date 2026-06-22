from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "FridayRun_Project_Documentation.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "666666"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
WARNING_FILL = "FFF5CC"
GREEN_FILL = "EAF6EF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="D9E2EC") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_after = Pt(4)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(14)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", 1)
        code.font.name = "Courier New"
        code._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        code.font.size = Pt(8.5)
        code.paragraph_format.space_after = Pt(4)
        code.paragraph_format.line_spacing = 1.0

    if "Callout" not in styles:
        callout = styles.add_style("Callout", 1)
        callout.font.name = "Calibri"
        callout.font.size = Pt(10.5)
        callout.paragraph_format.left_indent = Inches(0.15)
        callout.paragraph_format.right_indent = Inches(0.15)
        callout.paragraph_format.space_before = Pt(4)
        callout.paragraph_format.space_after = Pt(8)
        callout.paragraph_format.line_spacing = 1.2


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("FridayRun Project Documentation")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.add_run("FridayRun Project Documentation")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run("Architecture, operations, security, deployment, and user guide")
    meta = doc.add_table(rows=4, cols=2)
    set_table_geometry(meta, [2100, 7260])
    set_table_borders(meta)
    rows = [
        ("Application", "FridayRun"),
        ("Purpose", "Australian work-related travel records for tax deduction support"),
        ("Stack", "React + TypeScript, Tailwind CSS, Express, SQLite, Google Maps, Recharts"),
        ("Security model", "Multi-user accounts, HttpOnly cookie sessions, encrypted TFN storage, super-admin role"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        set_cell_shading(row.cells[0], HEADER_FILL)
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].text = value
    doc.add_paragraph()


def add_callout(doc: Document, title: str, text: str, fill=LIGHT_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, "C8D4E3")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.style = doc.styles["Callout"]
    run = p.add_run(f"{title}: ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p.add_run(text)
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, "D0D7DE")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F8FA")
    for idx, line in enumerate(text.strip("\n").splitlines()):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.style = doc.styles["Code Block"]
        p.add_run(line)
    doc.add_paragraph()


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(INK)
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    set_table_geometry(table, widths)
    doc.add_paragraph()


def toc(doc: Document) -> None:
    doc.add_heading("Contents", level=1)
    add_bullets(
        doc,
        [
            "Executive summary and feature inventory",
            "Architecture and codebase map",
            "User workflows and administrative workflows",
            "Database schema, API surface, and reporting model",
            "Security controls and TFN handling",
            "Local development and production deployment",
            "Cloudflare Tunnel and Apache reverse proxy operations",
            "Troubleshooting, maintenance, and future roadmap",
        ],
    )


def build() -> None:
    doc = Document()
    configure_styles(doc)
    add_footer(doc)
    add_title(doc)
    toc(doc)

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "FridayRun is a responsive web application for recording Australian work-related travel. "
        "It helps users maintain kilometre records for the Australian financial year, export tax-ready reports, "
        "and manage personal account details required on those reports."
    )
    add_callout(
        doc,
        "Current production posture",
        "The application supports multi-user authentication, super-admin user management, HttpOnly cookie sessions, TFN encryption at rest, Google Maps address support, and Cloudflare Tunnel self-hosting.",
        GREEN_FILL,
    )
    doc.add_heading("Primary capabilities", level=2)
    add_bullets(
        doc,
        [
            "Create, edit, delete, search, and filter travel entries.",
            "Use Google Places autocomplete for account addresses and trip destinations.",
            "Display trip locations on a map and calculate driving distance.",
            "Mark trips as one-way or return trips so reports are not misleading.",
            "Automatically calculate weekly, monthly, and Australian financial-year totals.",
            "Export CSV, PDF, and Excel reports with user details and trip type.",
            "Create accounts with full name, address, email, TFN, and strong password validation.",
            "Manage users from a super-admin dashboard.",
        ],
    )

    doc.add_heading("2. System Architecture", level=1)
    doc.add_paragraph(
        "FridayRun is implemented as a workspace-based full-stack TypeScript project. "
        "The production build serves the React application from the Express server, while local development runs Vite and Express separately."
    )
    add_table(
        doc,
        ["Layer", "Technology", "Purpose"],
        [
            ["Frontend", "React, TypeScript, Vite, Tailwind CSS, Recharts", "User interface, dashboards, forms, charts, report download workflow"],
            ["Maps", "Google Maps JavaScript API, Places API, Distance Matrix API", "Address autocomplete, map display, route distance calculation"],
            ["Backend", "Node.js, Express, TypeScript", "REST API, auth, validation, reports, static app serving"],
            ["Database", "SQLite via better-sqlite3", "Local/small deployment persistence for users, sessions, and trips"],
            ["Reports", "PDFKit, ExcelJS, CSV generation", "Downloadable tax log exports"],
            ["Deployment", "Apache reverse proxy, Cloudflare Tunnel, systemd or Docker", "Self-hosted production runtime"],
        ],
        [1450, 2600, 5310],
    )
    doc.add_heading("Repository layout", level=2)
    add_table(
        doc,
        ["Path", "Description"],
        [
            ["client/", "React frontend application."],
            ["client/src/components/", "Dashboard, forms, auth, admin, maps, and report UI components."],
            ["client/src/lib/", "API client and date helpers."],
            ["server/src/", "Express API, auth, validation, reports, repository, database, and security modules."],
            ["docs/self-hosting.md", "Apache and Cloudflare Tunnel deployment guide."],
            [".env.example", "Production and local environment variable template."],
            ["Dockerfile, docker-compose.yml", "Containerized local deployment support."],
        ],
        [2600, 6760],
    )

    doc.add_heading("3. User Workflows", level=1)
    doc.add_heading("Account creation", level=2)
    add_numbered(
        doc,
        [
            "User opens FridayRun and selects Create account.",
            "User enters full name, address, email, TFN, and a strong password.",
            "Address can be selected with Google Places autocomplete restricted to Australia.",
            "Password must be at least 12 characters and include uppercase, lowercase, number, and symbol characters.",
            "Server validates input, encrypts the TFN, hashes the password, creates the user, and sets an HttpOnly session cookie.",
        ],
    )
    doc.add_heading("Logging a trip", level=2)
    add_numbered(
        doc,
        [
            "User enters trip date, vehicle, optional odometer values, and optional notes.",
            "User selects start and end destinations with Google Places autocomplete or types manually.",
            "If Maps is configured, user can calculate the driving distance.",
            "If Round trip is ticked, the app records twice the one-way calculated distance.",
            "The trip is saved against the signed-in user's account and assigned to the correct Australian financial year.",
        ],
    )
    doc.add_heading("Reports", level=2)
    add_bullets(
        doc,
        [
            "Supported periods: custom date range, week, month, and financial year.",
            "Supported formats: CSV, PDF, and Excel (.xlsx).",
            "Reports include full name, address, email, TFN, trip date, start, end, trip type, kilometres, vehicle, notes, and total kilometres.",
            "Trip type is shown as One-way or Return trip.",
        ],
    )

    doc.add_heading("4. Administrative Workflows", level=1)
    doc.add_paragraph(
        "A super-admin account is seeded on first startup when no users exist. The seed values come from environment variables. "
        "After the account exists, changing .env does not change the stored admin credentials."
    )
    add_table(
        doc,
        ["Admin task", "Where", "Notes"],
        [
            ["View all users", "Admin tab", "Available only to role super_admin."],
            ["Change a user's role", "Admin user table", "Can promote/demote users between user and super_admin."],
            ["Delete users", "Admin user table", "Super-admin accounts are protected from delete by the current delete route."],
            ["Update own details", "Account tab", "Users can update full name, address, email, and TFN."],
            ["Change own password", "Account tab", "Requires current password and a new strong password."],
        ],
        [2300, 2300, 4760],
    )

    doc.add_heading("5. Database Model", level=1)
    doc.add_paragraph(
        "SQLite is currently used for local and small single-server deployments. For larger multi-user production workloads, PostgreSQL is recommended as the next persistence step."
    )
    add_table(
        doc,
        ["Table", "Purpose", "Important fields"],
        [
            ["users", "Account profiles and roles", "id, full_name, address, email, encrypted tfn, password_hash, role"],
            ["sessions", "Server-side login sessions", "token_hash, user_id, expires_at"],
            ["travel_entries", "User-owned trip records", "user_id, date, financial_year, start_location, end_location, kilometres, round_trip"],
        ],
        [1800, 2600, 4960],
    )
    doc.add_heading("Travel entry schema", level=2)
    add_code(
        doc,
        """
CREATE TABLE travel_entries (
  id TEXT PRIMARY KEY,
  user_id TEXT REFERENCES users(id) ON DELETE CASCADE,
  date TEXT NOT NULL,
  financial_year TEXT NOT NULL,
  start_location TEXT NOT NULL,
  end_location TEXT NOT NULL,
  kilometres REAL NOT NULL CHECK (kilometres >= 0),
  calculated_kilometres REAL,
  distance_source TEXT NOT NULL DEFAULT 'manual',
  notes TEXT,
  vehicle TEXT,
  odometer_start REAL,
  odometer_end REAL,
  round_trip INTEGER NOT NULL DEFAULT 0,
  created_at TEXT NOT NULL DEFAULT (datetime('now')),
  updated_at TEXT NOT NULL DEFAULT (datetime('now'))
);
""",
    )

    doc.add_heading("6. API Reference", level=1)
    add_table(
        doc,
        ["Endpoint", "Method", "Auth", "Purpose"],
        [
            ["/api/health", "GET", "No", "Health check."],
            ["/api/auth/signup", "POST", "No", "Create user and set HttpOnly session cookie."],
            ["/api/auth/login", "POST", "No", "Verify credentials and set HttpOnly session cookie."],
            ["/api/auth/logout", "POST", "Yes", "Delete server session and clear cookie."],
            ["/api/me", "GET", "Yes", "Return current user profile."],
            ["/api/me", "PUT", "Yes", "Update current user profile."],
            ["/api/me/password", "PUT", "Yes", "Change current user's password."],
            ["/api/admin/users", "GET", "Super admin", "List all users."],
            ["/api/admin/users/:id", "PUT", "Super admin", "Update user profile and role."],
            ["/api/admin/users/:id", "DELETE", "Super admin", "Delete non-super-admin user."],
            ["/api/dashboard", "GET", "Yes", "User-scoped kilometre totals and chart data."],
            ["/api/entries", "GET/POST", "Yes", "List or create user-scoped trip entries."],
            ["/api/entries/:id", "GET/PUT/DELETE", "Yes", "Read, update, or delete a user's own trip entry."],
            ["/api/financial-years", "GET", "Yes", "List financial years with entries for current user."],
            ["/api/reports", "GET", "Yes", "Download CSV, PDF, or Excel report."],
        ],
        [2650, 1000, 1350, 4360],
    )

    doc.add_heading("7. Security Model", level=1)
    add_callout(
        doc,
        "Sensitive data warning",
        "TFNs are sensitive personal information. Use HTTPS, restrict admin access, encrypt backups, and carefully control server/database access.",
        WARNING_FILL,
    )
    add_table(
        doc,
        ["Control", "Implementation"],
        [
            ["Password hashing", "Node crypto scrypt with per-password salt."],
            ["Password policy", "Minimum 12 characters with uppercase, lowercase, number, and symbol."],
            ["Session storage", "Server-side sessions table stores only SHA-256 token hashes."],
            ["Browser session handling", "HttpOnly cookies; JavaScript does not store session tokens."],
            ["TFN storage", "AES-256-GCM encryption with TFN_ENCRYPTION_KEY."],
            ["User scoping", "Trip queries and reports are filtered by authenticated user_id."],
            ["Admin access", "Role-based super_admin gate for user management routes."],
            ["CORS", "APP_ORIGIN allow-list with credentials enabled."],
        ],
        [2600, 6760],
    )
    doc.add_heading("Critical production secrets", level=2)
    add_bullets(
        doc,
        [
            "SESSION_SECRET: long random secret used as part of runtime security configuration.",
            "TFN_ENCRYPTION_KEY: 32-byte base64 or 64-character hex key; must remain stable to decrypt stored TFNs.",
            "SUPER_ADMIN_PASSWORD: only used when the first user is seeded.",
            "VITE_GOOGLE_MAPS_API_KEY: restrict in Google Cloud to the production domain and required APIs.",
        ],
    )

    doc.add_heading("8. Environment Variables", level=1)
    add_table(
        doc,
        ["Variable", "Example", "Purpose"],
        [
            ["NODE_ENV", "production", "Enables production static serving and secure defaults."],
            ["PORT", "3000", "Express listen port."],
            ["APP_ORIGIN", "https://fridayrun.example.com", "Allowed browser origin for CORS."],
            ["VITE_API_BASE_URL", "/api", "Frontend API base path; use /api in production."],
            ["SESSION_SECRET", "<random>", "Long random secret."],
            ["SESSION_COOKIE_SECURE", "true", "Use secure cookies behind HTTPS."],
            ["TFN_ENCRYPTION_KEY", "<32-byte base64>", "Encrypts and decrypts TFNs at rest."],
            ["DATABASE_URL", "./data/fridayrun.sqlite", "SQLite database path."],
            ["VITE_GOOGLE_MAPS_API_KEY", "<restricted key>", "Google Maps browser key."],
            ["SUPER_ADMIN_*", "various", "Initial super-admin seed values."],
        ],
        [2500, 2500, 4360],
    )
    doc.add_heading("Production .env template", level=2)
    add_code(
        doc,
        """
NODE_ENV=production
PORT=3000
APP_ORIGIN=https://fridayrun.example.com
VITE_API_BASE_URL=/api
SESSION_SECRET=<generated secret>
SESSION_COOKIE_SECURE=true
TFN_ENCRYPTION_KEY=<openssl rand -base64 32>
DATABASE_URL=./data/fridayrun.sqlite
VITE_GOOGLE_MAPS_API_KEY=<restricted Google Maps key>
SUPER_ADMIN_FULL_NAME=<admin name>
SUPER_ADMIN_ADDRESS=<admin address>
SUPER_ADMIN_EMAIL=<admin email>
SUPER_ADMIN_TFN=<admin TFN>
SUPER_ADMIN_PASSWORD=<strong password>
""",
    )

    doc.add_heading("9. Local Development", level=1)
    add_numbered(
        doc,
        [
            "Copy .env.example to .env and fill required values.",
            "Run npm install from the repository root.",
            "Run npm run dev to start Vite and Express concurrently.",
            "Open the Vite frontend on http://localhost:5173.",
            "Use http://localhost:3000/api/health to verify the API.",
        ],
    )
    add_code(
        doc,
        """
cp .env.example .env
npm install
npm run dev
curl http://localhost:3000/api/health
""",
    )

    doc.add_heading("10. Production Self-hosting", level=1)
    doc.add_paragraph(
        "The recommended self-hosted topology is Cloudflare Tunnel to Apache reverse proxy to the Node application. "
        "The origin server should not expose public inbound HTTP/HTTPS ports."
    )
    add_code(
        doc,
        """
Cloudflare Tunnel -> Apache reverse proxy -> FridayRun Node app -> SQLite database
""",
    )
    doc.add_heading("Build and service", level=2)
    add_code(
        doc,
        """
cd /opt/fridayrun
npm install
npm run build
sudo systemctl restart fridayrun
""",
    )
    add_table(
        doc,
        ["File", "Location", "Purpose"],
        [
            ["Application code", "/opt/fridayrun", "Project root, build output, server, and frontend."],
            ["Environment file", "/opt/fridayrun/.env", "Production runtime configuration."],
            ["Apache site config", "/etc/apache2/sites-available/fridayrun.conf", "Reverse proxy vhost."],
            ["Enabled Apache site", "/etc/apache2/sites-enabled/", "Symlink created by a2ensite."],
            ["Cloudflared config", "/etc/cloudflared/config.yml", "Tunnel ingress rules."],
        ],
        [2300, 3000, 4060],
    )
    doc.add_heading("systemd service", level=2)
    add_code(
        doc,
        """
[Unit]
Description=FridayRun
After=network.target

[Service]
WorkingDirectory=/opt/fridayrun
ExecStart=/usr/bin/npm run start --workspace server
Restart=always
Environment=NODE_ENV=production

[Install]
WantedBy=multi-user.target
""",
    )

    doc.add_heading("11. Apache and Cloudflare Tunnel", level=1)
    doc.add_heading("Apache port and virtual host", level=2)
    add_code(
        doc,
        """
# /etc/apache2/ports.conf
Listen 127.0.0.1:8080
""",
    )
    add_code(
        doc,
        """
<VirtualHost 127.0.0.1:8080>
  ServerName fridayrun.example.com

  ProxyPreserveHost On
  ProxyPass / http://127.0.0.1:3000/
  ProxyPassReverse / http://127.0.0.1:3000/

  RequestHeader set X-Forwarded-Proto "https"
  Header always set X-Content-Type-Options "nosniff"
  Header always set Referrer-Policy "no-referrer"
</VirtualHost>
""",
    )
    doc.add_heading("Cloudflare Tunnel ingress", level=2)
    add_code(
        doc,
        """
tunnel: fridayrun
credentials-file: /etc/cloudflared/fridayrun.json

ingress:
  - hostname: fridayrun.example.com
    service: http://127.0.0.1:8080
  - service: http_status:404
""",
    )
    doc.add_heading("Verification chain", level=2)
    add_code(
        doc,
        """
curl http://127.0.0.1:3000/api/health
curl http://127.0.0.1:8080/api/health
cloudflared tunnel run fridayrun
""",
    )

    doc.add_heading("12. Google Maps Configuration", level=1)
    add_bullets(
        doc,
        [
            "Enable Maps JavaScript API.",
            "Enable Places API.",
            "Enable Distance Matrix API for the current implementation.",
            "Restrict the browser key to the production hostname.",
            "Keep VITE_GOOGLE_MAPS_API_KEY in .env and rebuild after changes.",
        ],
    )

    doc.add_heading("13. Troubleshooting", level=1)
    add_table(
        doc,
        ["Symptom", "Likely cause", "Fix"],
        [
            ["Origin not allowed", "APP_ORIGIN does not match browser origin", "Set APP_ORIGIN to the exact https:// hostname and restart."],
            ["Invalid email or password after .env change", "Admin account was already seeded", "Use original login, reset password in DB/script, or reset database."],
            ["Cloudflared connect refused 127.0.0.1:8080", "Apache is not listening on 8080", "Add Listen 127.0.0.1:8080, enable vhost, reload Apache."],
            ["Browser calls localhost", "VITE_API_BASE_URL built as localhost", "Set VITE_API_BASE_URL=/api and rebuild."],
            ["Secure cookie not set on IP/HTTP test", "SESSION_COOKIE_SECURE=true over HTTP", "Temporarily set false only for local HTTP testing."],
            ["TFNs unreadable after key change", "TFN_ENCRYPTION_KEY changed", "Restore original key from secure backup."],
        ],
        [2300, 3000, 4060],
    )

    doc.add_heading("14. Maintenance and Operations", level=1)
    add_bullets(
        doc,
        [
            "Back up the SQLite database and the TFN_ENCRYPTION_KEY together, but store them securely and separately where possible.",
            "Do not commit .env, SQLite databases, or generated WAL/SHM files.",
            "Patch Node.js, Apache, cloudflared, npm packages, and the Linux server regularly.",
            "Use Cloudflare Access or another MFA-protected layer for admin access if possible.",
            "Monitor systemd service health with systemctl status fridayrun and journalctl -u fridayrun.",
            "Use npm run typecheck and npm run build before deployment.",
        ],
    )

    doc.add_heading("15. Future Roadmap", level=1)
    add_table(
        doc,
        ["Priority", "Enhancement", "Reason"],
        [
            ["High", "PostgreSQL migration", "Better multi-user production durability, backup tooling, concurrency, and hosting options."],
            ["High", "MFA for super admins", "Additional protection for sensitive TFN-bearing system."],
            ["Medium", "Formal audit log", "Track admin changes, report downloads, and profile updates."],
            ["Medium", "Cloudflare Access integration guide", "Make deployment safer for private/internal usage."],
            ["Medium", "Automated backup/restore scripts", "Reduce operational risk."],
            ["Low", "PWA polish", "Improve mobile home-screen experience."],
        ],
        [1200, 3000, 5160],
    )

    doc.add_heading("Appendix A. Commands Quick Reference", level=1)
    add_code(
        doc,
        """
npm run typecheck
npm run build
sudo systemctl status fridayrun
sudo journalctl -u fridayrun -f
sudo apache2ctl configtest
sudo systemctl reload apache2
sudo ss -ltnp | grep 8080
curl http://127.0.0.1:3000/api/health
curl http://127.0.0.1:8080/api/health
sudo journalctl -xeu cloudflared.service
""",
    )

    doc.add_heading("Appendix B. Compliance Notes", level=1)
    add_callout(
        doc,
        "Not legal or tax advice",
        "FridayRun helps maintain travel records, but users should confirm ATO evidence requirements and privacy obligations with appropriate professional advice.",
        WARNING_FILL,
    )
    add_bullets(
        doc,
        [
            "Australian financial year is calculated as 1 July to 30 June.",
            "Reports include user profile details so exported records identify the taxpayer.",
            "Round trips are explicitly labeled to avoid misleading trip direction or distance assumptions.",
            "TFN collection increases privacy obligations; remove or make TFN optional if not genuinely required.",
        ],
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
